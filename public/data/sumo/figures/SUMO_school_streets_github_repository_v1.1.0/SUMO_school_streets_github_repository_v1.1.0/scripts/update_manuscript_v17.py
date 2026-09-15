#!/usr/bin/env python3
"""Update the publication-support manuscript to release 1.1.0 / SUMO 1.27.1."""

from __future__ import annotations

import argparse
import csv
from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Emu


ROOT = Path(__file__).resolve().parents[1]
SUMMARY = ROOT / "analysis" / "replicate_summary.csv"
FIGURE_6 = ROOT / "docs" / "figures" / "manuscript_figure6_primary_results_v1.1.png"
FIGURE_7 = ROOT / "docs" / "figures" / "manuscript_figure7_supplementary_results_v1.1.png"
FIGURE_1 = ROOT / "docs" / "figures" / "manuscript_figure1_workflow_v1.1.png"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    return parser.parse_args()


def load_summary() -> dict[str, dict[str, float]]:
    with SUMMARY.open(newline="", encoding="utf-8") as handle:
        result = {}
        for row in csv.DictReader(handle):
            result[row["scenario"]] = {
                key: float(value)
                for key, value in row.items()
                if key != "scenario" and value not in ("", None) and key != "interpretation_status"
            }
        return result


def set_paragraph(doc: Document, prefix: str, text: str) -> None:
    matches = [paragraph for paragraph in doc.paragraphs if paragraph.text.startswith(prefix)]
    if len(matches) != 1:
        raise ValueError(f"Expected one paragraph beginning {prefix!r}; found {len(matches)}")
    matches[0].text = text


def set_cell(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    runs = paragraph.runs
    if runs:
        runs[0].text = text
        for run in runs[1:]:
            run._element.getparent().remove(run._element)
    else:
        paragraph.add_run(text)
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)


def replace_picture(paragraph, image_path: Path, alt_text: str) -> None:
    inline = paragraph._p.xpath(".//wp:inline")
    if len(inline) != 1:
        raise ValueError("Expected one inline picture in replacement paragraph")
    extent = inline[0].find(qn("wp:extent"))
    width = int(extent.get("cx"))
    height = int(extent.get("cy"))
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Emu(width), height=Emu(height))
    doc_properties = run._r.xpath(".//wp:docPr")
    if len(doc_properties) != 1:
        raise ValueError("Expected one drawing properties element after picture replacement")
    doc_properties[0].set("descr", alt_text)
    doc_properties[0].set("title", alt_text)


def fmt(row: dict[str, float], metric: str, digits: int, sd_digits: int | None = None) -> str:
    sd_digits = digits if sd_digits is None else sd_digits
    return f"{row[f'{metric}_mean_valid']:.{digits}f} ± {row[f'{metric}_sd_valid']:.{sd_digits}f}"


def update_tables(doc: Document, rows: dict[str, dict[str, float]]) -> None:
    # Table 2: calibration target.
    set_cell(doc.tables[1].rows[7].cells[2], "Five-seed baseline: 36.34% ± 0.59 percentage points")

    # Table 6: principal full-demand comparison.
    primary = (("current", "Baseline"), ("A_OBS", "A"), ("B_OBS", "B"), ("C_OBS", "C0"))
    table = doc.tables[5]
    for table_row, (scenario, label) in zip(table.rows[1:], primary):
        row = rows[scenario]
        values = (
            label,
            fmt(row, "frontage_spot_motorized_mean_speed_kmh", 2),
            fmt(row, "frontage_spot_motorized_over_36_pct", 2),
            fmt(row, "mean_time_loss_s", 2),
            fmt(row, "CO2_kg", 1),
            fmt(row, "CO2_g_per_vkm", 1),
            fmt(row, "bus_mean_time_loss_s", 2),
        )
        for cell, value in zip(table_row.cells, values):
            set_cell(cell, value)

    # Table 7: A/B/C0 demand-response sensitivities.
    table = doc.tables[6]
    cases = (
        ("A_OBS", "A", "OBS"), ("A_DR15", "A", "DR15"), ("A_DR30", "A", "DR30"),
        ("B_OBS", "B", "OBS"), ("B_DR15", "B", "DR15"), ("B_DR30", "B", "DR30"),
        ("C_OBS", "C0", "OBS"), ("C_DR15", "C0", "DR15"), ("C_DR30", "C0", "DR30"),
    )
    for table_row, (scenario, design, demand) in zip(table.rows[1:], cases):
        row = rows[scenario]
        values = (
            design,
            demand,
            f"{row['frontage_spot_motorized_passages_mean_valid']:,.0f}",
            f"{row['frontage_spot_motorized_mean_speed_kmh_mean_valid']:.2f}",
            f"{row['frontage_spot_motorized_over_36_pct_mean_valid']:.2f}",
            f"{row['CO2_kg_mean_valid']:.1f}",
        )
        for cell, value in zip(table_row.cells, values):
            set_cell(cell, value)

    # Table 8: C1 matched-control changes.
    table = doc.tables[7]
    x0 = rows["C1_X0"]
    c1_cases = (
        ("C1_X0", "0% (0)"),
        ("C1_L25", "25% (382)"),
        ("C1_M50", "50% (764)"),
        ("C1_H75", "75% (1,146)"),
    )
    for table_row, (scenario, share) in zip(table.rows[1:], c1_cases):
        row = rows[scenario]
        if scenario == "C1_X0":
            values = (scenario, share, "0.0%", "0.0%", "0.00", "—", "0.0")
        else:
            speed = 100 * (row["frontage_spot_motorized_mean_speed_kmh_mean_valid"] / x0["frontage_spot_motorized_mean_speed_kmh_mean_valid"] - 1)
            co2 = 100 * (row["CO2_g_per_vkm_mean_valid"] / x0["CO2_g_per_vkm_mean_valid"] - 1)
            bus = row["bus_mean_time_loss_s_mean_valid"] - x0["bus_mean_time_loss_s_mean_valid"]
            values = (
                scenario,
                share,
                f"{speed:+.1f}%".replace("-", "−"),
                f"{co2:+.1f}%",
                f"{bus:+.2f}",
                f"{row['pedestrian_crossing_mean_waiting_s_mean_valid']:.2f}",
                f"{row['collision_output_intermodal_mean_valid']:.1f}",
            )
        for cell, value in zip(table_row.cells, values):
            set_cell(cell, value)

    # Table 9: one-way frontage-access experiments.
    table = doc.tables[8]
    one_way = (
        ("current", "Baseline", "—", "Reference"),
        ("OW_FULL_OBS", "D / OW_FULL", "−49.4%", "Full one-way frontage"),
        ("OW_REDUCED_OBS", "E / OW_REDUCED", "−49.4%", "Reduced one-way frontage"),
    )
    for table_row, (scenario, label, change, scope) in zip(table.rows[1:], one_way):
        row = rows[scenario]
        values = (
            label,
            f"{row['frontage_spot_motorized_passages_mean_valid']:,.0f}",
            change,
            f"{row['frontage_spot_motorized_mean_speed_kmh_mean_valid']:.2f}",
            f"{row['frontage_spot_motorized_over_36_pct_mean_valid']:.2f}",
            f"{row['bus_mean_time_loss_s_mean_valid']:.2f}",
            scope,
        )
        for cell, value in zip(table_row.cells, values):
            set_cell(cell, value)


def update_text(doc: Document) -> None:
    set_paragraph(doc, "School-street redesigns seek", "School-street redesigns seek to reduce road danger and improve public space, yet their operational and emissions effects are difficult to quantify before implementation. This study presents an evidence-aligned SUMO 1.27.1 workflow that translates an expert road-safety inspection into reproducible microsimulation cases for the De l’Autre Côté de l’École corridor in Brussels. A 24 h baseline was constructed from Telraam 15-min counts, a near-date Line 34 timetable proxy, OpenStreetMap geometry and Belgian fleet data. The model reproduces 7,597 cars, 803 large vehicles, 965 two-wheelers and 1,528 pedestrians; five seeds yield 36.34% of motorized frontage passages above 36 km/h against a 37% report target. At full observed demand, relocating both bus stops (Scenario A’s quantifiable component) increased mean frontage speed from 32.77 to 35.27 km/h, whereas the chicane proxy (B) and 20 km/h operational proxy (C0) reduced it to 29.27 and 22.35 km/h. Their shares above 36 km/h were 11.59% and 0.09%, respectively. A matched pedestrian-priority sensitivity reduced speed by 3.3–10.1% as the assumed crossing share rose from 25% to 75%, but increased CO₂/VKT by 0.6–1.8% and produced more intermodal overlap diagnostics. Two supplementary one-way concepts reduced modeled frontage passages by 49.4% but left the speeding share above the baseline; onward diversion was outside the calibrated boundary. The resulting version-pinned repository distinguishes modeled features, proxies, sensitivities and unmodeled requirements, providing a transparent basis for pre-implementation planning.")
    set_paragraph(doc, "A common desired-speed factor distribution", "A common desired-speed factor distribution, normc(1.372, 0.20, 0.75, 2.00), was fitted to the report’s secondary aggregate target that 37% of motorized vehicles exceeded 36 km/h. Across seeds 42–46, the SUMO 1.27.1 baseline produces 36.34%, an error of −0.66 percentage points. The same driver population and 0.8 m sublane resolution are retained in every design case.")
    set_paragraph(doc, "The complete experiment contains 20 cases.", "The complete experiment contains 20 cases. Baseline, A/B/C0 at observed demand, their DR15/DR30 variants, four C1 cases, and D/E with observed and reduced demand were each run for 86,400 s using SUMO 1.27.1 and seeds 42–46. A replicate contributes to reported means only if every loaded vehicle completes; adopted observed-demand cases must have no vehicle or person teleports. D/E additionally require zero collision diagnostics. Release 1.1.0 / model v21.2 is version-pinned to this runtime.")
    set_paragraph(doc, "The final package passed", "The final package passed 307 automated integrity checks with 10 disclosed evidence warnings and no failed checks. Checks cover deterministic ID sets, demand and timetable totals, route ODs, permissions, C1 topology, common behavior and fleet assignment, five-seed completion, analysis consistency and manifest integrity. The warnings retain the substantive evidence boundaries: C1 geometry, near-date timetable, absent parking supply, D/E diversion, directional uncertainty, detector-time redistribution, fleet transfer, C1 overlaps, NOx sensitivity and A’s qualitative visual scope. One first C_OBS seed-44 attempt produced incomplete XML; the existing integrity safeguard rejected it, reran the complete seed, and admitted only the clean rerun to the processed tables.")
    set_paragraph(doc, "All five baseline replicates", "All five baseline replicates completed 9,365 vehicles and 1,528 pedestrians without teleports. The school-frontage detector recorded 7,597 cars, 803 large vehicles and 965 two-wheelers daily, matching the observed totals. The five-seed motorized spot speed was 32.77 ± 0.15 km/h, and 36.34% ± 0.59 percentage points of passages exceeded 36 km/h. The latter differs from the 37% target by −0.66 percentage points. Input-bin counts are exact, although travel and queuing shift some detector passages into adjacent 15-min bins.")
    set_paragraph(doc, "Scenario A’s quantified component", "Scenario A’s quantified component—the relocation of both stops—raises mean frontage speed by 7.64% to 35.27 km/h and raises the share above 36 km/h by 12.59 percentage points. Mean route time loss falls from 20.86 to 5.64 s and CO₂/VKT falls by 21.75%, while bus time loss rises from 6.08 to 9.39 s under the SUMO 1.27.1 definition. This does not evaluate the full enhanced-zone design: lighting, markings and sightline measures remain qualitative. The result instead shows that stop relocation alone is not a speed-calming mechanism.")
    set_paragraph(doc, "Scenario B reduces", "Scenario B reduces mean spot speed by 10.67% to 29.27 km/h and the share above 36 km/h by 24.75 percentage points to 11.59%. Scenario C0 produces the strongest operational speed effect: 22.35 km/h (−31.78%) and only 0.09% above 36 km/h. Under the implemented full-demand proxies, CO₂/VKT is 18.00% lower than the baseline in B and 13.34% lower in C0. These are route-based emissions results for the modeled designs, not evidence about concentrations, exposure, or the unmodeled public-realm elements.")
    set_paragraph(doc, "Relative to each design’s OBS case", "Relative to each design’s OBS case, DR15 lowers total CO₂ by 13.7–14.2% and DR30 by 26.4–27.3%, broadly following the reduction in represented VKT. Mean frontage speed, however, rises by approximately 1.6–4.7% as demand falls. In A, the above-36 km/h share increases from 48.93% to 55.17% at DR30. Demand reduction can therefore lower absolute emissions while weakening speed discipline; it is not a substitute for a physical speed-management treatment.")
    set_paragraph(doc, "The response is systematic", "The response is systematic: greater assumed crossing activity lowers vehicle speed, adds 0.27–0.92 s of mean bus time loss, and keeps mean crossing wait near 0.65 s, but also increases braking and CO₂/VKT. Intermodal overlap diagnostics range from 5.6 to 10.8 per run; many involve bicycles. They are not field-calibrated conflicts or predicted crashes, so C1 supports an operational sensitivity conclusion only.")
    set_paragraph(doc, "Both adopted one-way formulations", "Both adopted one-way formulations complete all five seeds without teleports or collision diagnostics. The 49.4% frontage reduction follows the balanced direction assumption and retained two-way buses; it is not an observed diversion forecast. Neither option calms the remaining traffic: the share above 36 km/h increases by approximately 3.6 percentage points relative to the baseline. Route-based time, VKT and emissions after the boundary are therefore not compared.")
    set_paragraph(doc, "The corrected full-demand comparison", "The corrected full-demand comparison shows that A–C0 do not require an assumed traffic reduction to operate: all 9,365 vehicles and 1,528 pedestrians complete in every replicate. Scenario A then isolates a useful negative finding. Relocating the bus stops reduces overall route time loss and tailpipe CO₂ intensity, but under SUMO 1.27.1 it raises mean bus time loss from 6.08 to 9.39 s and also increases school-frontage speed and speeding prevalence. The visual, lighting and visibility measures proposed by the report may still improve recognition and behavior, but the available evidence does not justify assigning them a numerical response coefficient. Stop relocation alone should not be presented as the modeled cause of safer speed behavior.")
    set_paragraph(doc, "The methodological contribution is", "The methodological contribution is the auditable evidence chain. The repository records which source defined each requirement, whether that requirement is modeled, proxied, varied or excluded, and which comparisons remain valid. It includes immutable templates, preparation and execution scripts, five-seed processed outputs, validation tables, model-derived schematics, a guidebook and backend/frontend handovers. Release 1.1.0 additionally provides a machine-readable 1.24.0-to-1.27.1 comparison, so runtime-dependent time-loss, emissions and interaction values cannot be silently mixed. This structure reduces the risk that a visually plausible scenario is mistaken for a calibrated representation of every design feature.")
    set_paragraph(doc, "Fleet uncertainty is likewise asymmetric", "Fleet uncertainty is likewise asymmetric. Replacing the newer passenger Euro 6d proxy with a conservative Euro 5 proxy changes CO₂ by only about 1.0–2.0% across the tested cases but increases absolute NOx by approximately 314–498%. The principal CO₂ patterns are therefore more stable than the absolute NOx inventory. Locally observed Euro-stage and vehicle-age data would be required for stronger pollutant-specific conclusions [65]–[70].")
    set_paragraph(doc, "The evidence remains limited", "The evidence remains limited by one observed day, no directional or OD counts, fixed corridor-boundary routes, and a near-date rather than exact-date bus timetable. The model does not include passenger boarding, parking occupancy or search, a numerical response to lighting and markings, a continuous curbless surface, cross-anywhere pedestrian behavior, or calibrated yielding and conflict trajectories. C1 overlap events are simulation diagnostics, not crashes; D/E exclude onward diversion; and all emissions remain sensitive to fleet, driving-dynamics and SUMO-runtime assumptions. Because the 1.27.1 migration changed time-loss, emissions and some interaction diagnostics, publication values must remain version-pinned.")
    set_paragraph(doc, "This study developed an evidence-aligned", "This study developed an evidence-aligned, sensor-calibrated SUMO 1.27.1 workflow for pre-implementation assessment of school-street designs. Telraam 15-min counts, a near-date Line 34 timetable, OSM geometry, Belgian fleet data and expert design sources were combined in a 20-case, five-seed experiment. The baseline reproduces the observed daily multimodal totals and the report’s aggregate speed-exceedance target, while the accompanying version-pinned repository makes every scenario translation and inference boundary explicit.")
    set_paragraph(doc, "At full observed demand, Scenario A’s", "At full observed demand, Scenario A’s quantified bus-stop relocation reduced overall route time loss and CO₂ intensity but increased frontage speed, speeding prevalence and mean bus time loss; it should not stand in for the report’s unmodeled visual-calming package. The B chicane and C0 20 km/h proxy both improved speed discipline, reducing mean spot speed to 29.27 and 22.35 km/h and the share above 36 km/h to 11.59% and 0.09%, respectively. These findings are operational, not crash or exposure predictions.")
    set_paragraph(doc, "A publication-support release has been prepared", "Publication-support release 1.1.0 / model v21.2 has been prepared in two forms: a GitHub-ready repository with model templates, scripts, SUMO 1.27.1 processed results, guidebook and handovers, and a complete archival package containing the canonical generated cases and outputs. The release includes a machine-readable comparison with the historical SUMO 1.24.0 results. The persistent GitHub URL and Zenodo DOI will be inserted before publication. The project-supplied road-safety report and mobility presentation are referenced through a source register and SHA-256 checksums but are not redistributed pending confirmation of publication rights.")
    set_paragraph(doc, "Table 6. Primary full-demand results", "Table 6. Primary full-demand results under SUMO 1.27.1 (mean ± sample SD across seeds 42–46).")
    set_paragraph(doc, "Figure 6. Primary scenarios", "Figure 6. Primary scenarios at full observed demand under SUMO 1.27.1. Bars show five-seed means and error bars show sample SD. A represents bus-stop relocation only; B and C0 represent the adopted chicane and 20 km/h operational proxies.")
    set_paragraph(doc, "Figure 7. Bounded supplementary evidence", "Figure 7. Bounded supplementary evidence under SUMO 1.27.1. (a) C1 speed and CO₂/VKT changes are calculated against C1_X0. (b) D/E passage and speed-exceedance values describe the school frontage; off-network diversion is excluded.")


def update_metadata(doc: Document) -> None:
    properties = doc.core_properties
    properties.title = "Sensor-calibrated microsimulation of school-street interventions: mobility and emissions trade-offs in a Brussels living lab"
    properties.subject = "Updated evidence-aligned manuscript using repository release 1.1.0 / model v21.2 / SUMO 1.27.1"
    properties.author = "Filippos Lygerakis"
    properties.last_modified_by = "Filippos Lygerakis"
    properties.modified = datetime(2026, 8, 14, 12, 0, tzinfo=timezone.utc)
    properties.revision = 8
    properties.comments = "Clean manuscript update based on the validated SUMO 1.27.1 publication-support repository."


def main() -> None:
    args = parse_args()
    rows = load_summary()
    doc = Document(args.input)
    update_text(doc)
    update_tables(doc, rows)

    picture_paragraphs = [paragraph for paragraph in doc.paragraphs if paragraph._p.xpath(".//w:drawing")]
    image_names = []
    for paragraph in picture_paragraphs:
        blips = paragraph._p.xpath(".//a:blip")
        if len(blips) != 1:
            image_names.append("")
            continue
        rel_id = blips[0].get(qn("r:embed"))
        image_names.append(Path(paragraph.part.related_parts[rel_id].partname).name)
    mapping = dict(zip(image_names, picture_paragraphs))
    replace_picture(
        mapping["image9.png"],
        FIGURE_1,
        "Five-stage evidence-aligned SUMO workflow from source evidence and calibrated reference through evidence tiers, execution, quality assurance and decision outputs.",
    )
    replace_picture(
        mapping["image11.png"],
        FIGURE_6,
        "Three-panel bar chart comparing baseline and Scenarios A, B and C0 for mean school-frontage spot speed, share above 36 kilometres per hour and tailpipe carbon-dioxide intensity.",
    )
    replace_picture(
        mapping["image12.png"],
        FIGURE_7,
        "Two-panel supplementary results figure showing C1 speed and carbon-dioxide intensity changes by pedestrian crossing share, and one-way Scenario D and E frontage passages with speed exceedance.",
    )

    update_metadata(doc)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)
    print(f"Wrote {args.output}")


if __name__ == "__main__":
    main()
