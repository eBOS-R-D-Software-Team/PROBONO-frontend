#!/usr/bin/env python3
"""Build the reviewer-facing DOCX guidebook from docs/GUIDEBOOK.md."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "GUIDEBOOK.md"
OUTPUT = ROOT / "docs" / "SUMO_School_Streets_Guidebook_v1.1.docx"

# compact_reference_guide preset, resolved to explicit values.
PAGE_W = 12240
PAGE_H = 15840
MARGIN = 1440
CONTENT_W = 9360
TABLE_INDENT = 120
CELL_MARGINS = {"top": 40, "bottom": 40, "start": 120, "end": 120}

FONT = "Calibri"
INK = "263238"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "66737D"
TABLE_FILL = "E8EEF5"
LIGHT_FILL = "F4F6F9"
CAUTION_FILL = "FFF5D9"
BORDER = "C9D2DA"


def set_cell_margins(cell) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in CELL_MARGINS.items():
        elem = tc_mar.find(qn(f"w:{key}"))
        if elem is None:
            elem = OxmlElement(f"w:{key}")
            tc_mar.append(elem)
        elem.set(qn("w:w"), str(value))
        elem.set(qn("w:type"), "dxa")


def set_run(run, size=11, bold=None, italic=None, color=INK, font=FONT) -> None:
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_spacing(paragraph, before=0, after=6, line=1.25) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_hyperlink(paragraph, text: str, url: str):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), FONT)
    r_fonts.set(qn("w:hAnsi"), FONT)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "22")
    r_pr.extend([r_fonts, color, underline, size])
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.extend([r_pr, text_node])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


INLINE = re.compile(r"(\*\*.+?\*\*|`.+?`|\[[^\]]+\]\([^)]+\)|<https?://[^>]+>)")


def add_inline(paragraph, text: str, size=11, color=INK) -> None:
    text = text.replace("\\*", "*")
    position = 0
    for match in INLINE.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position:match.start()])
            set_run(run, size=size, color=color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run(run, size=size, bold=True, color=color)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run(run, size=max(size - 0.5, 8), color=DARK_BLUE, font="Consolas")
        elif token.startswith("["):
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token).groups()
            add_hyperlink(paragraph, label, url)
        else:
            url = token[1:-1]
            add_hyperlink(paragraph, url, url)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_run(run, size=size, color=color)


def paragraph_shading(paragraph, fill: str, border: str = BORDER) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), fill)
    p_pr.append(shade)
    p_bdr = OxmlElement("w:pBdr")
    for side in ("top", "left", "bottom", "right"):
        elem = OxmlElement(f"w:{side}")
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), "4")
        elem.set(qn("w:space"), "6")
        elem.set(qn("w:color"), border)
        p_bdr.append(elem)
    p_pr.append(p_bdr)


def set_style(style, size, color, before, after, line=1.25, bold=False) -> None:
    style.font.name = FONT
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold
    fmt = style.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_with_next = True


def setup_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    set_style(normal, 11, INK, 0, 6, 1.25, False)
    normal.paragraph_format.keep_with_next = False
    set_style(doc.styles["Heading 1"], 16, BLUE, 18, 10, 1.0, True)
    set_style(doc.styles["Heading 2"], 13, BLUE, 14, 7, 1.0, True)
    set_style(doc.styles["Heading 3"], 12, DARK_BLUE, 10, 5, 1.0, True)
    caption = doc.styles["Caption"]
    set_style(caption, 9, MUTED, 4, 8, 1.0, False)
    caption.font.italic = True


def numbering(doc: Document) -> tuple[int, int]:
    numbering_part = doc.part.numbering_part
    root = numbering_part.element
    abstract_ids = [int(e.get(qn("w:abstractNumId"))) for e in root.findall(qn("w:abstractNum"))]
    num_ids = [int(e.get(qn("w:numId"))) for e in root.findall(qn("w:num"))]
    next_abstract = max(abstract_ids or [0]) + 1
    next_num = max(num_ids or [0]) + 1

    def add_definition(abstract_id: int, num_id: int, fmt: str, text: str, font: str | None = None):
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), text)
        suff = OxmlElement("w:suff")
        suff.set(qn("w:val"), "tab")
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "540")
        tabs.append(tab)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "540")
        ind.set(qn("w:hanging"), "270")
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "80")
        spacing.set(qn("w:line"), "300")
        spacing.set(qn("w:lineRule"), "auto")
        p_pr.extend([tabs, ind, spacing])
        lvl.extend([start, num_fmt, lvl_text, suff, p_pr])
        if font:
            r_pr = OxmlElement("w:rPr")
            r_fonts = OxmlElement("w:rFonts")
            r_fonts.set(qn("w:ascii"), font)
            r_fonts.set(qn("w:hAnsi"), font)
            r_pr.append(r_fonts)
            lvl.append(r_pr)
        abstract.append(lvl)
        root.append(abstract)
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract_ref = OxmlElement("w:abstractNumId")
        abstract_ref.set(qn("w:val"), str(abstract_id))
        num.append(abstract_ref)
        root.append(num)

    add_definition(next_abstract, next_num, "bullet", "•", FONT)
    add_definition(next_abstract + 1, next_num + 1, "decimal", "%1.")
    return next_num, next_num + 1


def apply_num(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num])
    p_pr.append(num_pr)


def fixed_table_geometry(table, widths: list[int]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    for tag in ("tblW", "tblInd", "tblLayout"):
        existing = tbl_pr.find(qn(f"w:{tag}"))
        if existing is not None:
            tbl_pr.remove(existing)
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.extend([tbl_w, tbl_ind, layout])

    old_grid = table._tbl.tblGrid
    for child in list(old_grid):
        old_grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        old_grid.append(col)

    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def table_widths(rows: list[list[str]]) -> list[int]:
    columns = len(rows[0])
    scores = []
    for index in range(columns):
        values = [row[index] if index < len(row) else "" for row in rows]
        max_len = max(len(re.sub(r"[`*]", "", value)) for value in values)
        scores.append(max(7, min(max_len, 45)))
    min_width = 700 if columns >= 6 else 900
    widths = [max(min_width, int(CONTENT_W * score / sum(scores))) for score in scores]
    delta = CONTENT_W - sum(widths)
    widths[-1] += delta
    if widths[-1] < min_width:
        deficit = min_width - widths[-1]
        widths[-1] = min_width
        donors = sorted(range(columns - 1), key=lambda i: widths[i], reverse=True)
        for donor in donors:
            available = max(0, widths[donor] - min_width)
            take = min(available, deficit)
            widths[donor] -= take
            deficit -= take
            if deficit == 0:
                break
    return widths


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    columns = max(len(row) for row in rows)
    rows = [row + [""] * (columns - len(row)) for row in rows]
    table = doc.add_table(rows=len(rows), cols=columns)
    headers = [re.sub(r"[`*]", "", value).strip().lower() for value in rows[0]]
    if headers == ["design family", "full-demand case", "additional cases", "evidential role"]:
        widths = [1100, 1800, 2600, 3860]
    elif headers == [
        "case",
        "valid seeds",
        "frontage passages",
        "spot speed",
        "above 36 km/h",
        "co₂",
        "interpretation boundary",
    ]:
        widths = [900, 900, 1350, 1100, 1150, 900, 3060]
    else:
        widths = table_widths(rows)
    fixed_table_geometry(table, widths)
    table.style = "Table Grid"
    for r_index, (row, values) in enumerate(zip(table.rows, rows)):
        if r_index == 0:
            tr_pr = row._tr.get_or_add_trPr()
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            tr_pr.append(repeat)
        for cell, value in zip(row.cells, values):
            cell.text = ""
            shade = cell._tc.get_or_add_tcPr().find(qn("w:shd"))
            if r_index == 0:
                if shade is None:
                    shade = OxmlElement("w:shd")
                    cell._tc.get_or_add_tcPr().append(shade)
                shade.set(qn("w:fill"), TABLE_FILL)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if columns >= 4 and len(value) < 18 else WD_ALIGN_PARAGRAPH.LEFT
            set_paragraph_spacing(p, 0, 2, 1.08)
            add_inline(p, value, size=8.5 if columns >= 6 else 9)
            for run in p.runs:
                if r_index == 0:
                    run.bold = True
    spacer = doc.add_paragraph()
    set_paragraph_spacing(spacer, 0, 3, 1.0)


def add_image(doc: Document, path: Path, alt: str, figure_number: int) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    shape = run.add_picture(str(path), width=Inches(6.45))
    doc_pr = shape._inline.docPr
    doc_pr.set("descr", alt)
    caption = doc.add_paragraph(style="Caption")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(caption, f"Figure {figure_number}. {alt}", size=9, color=MUTED)


def page_field(paragraph) -> None:
    run = paragraph.add_run("Page ")
    set_run(run, size=9, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def setup_page(section) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_spacing(p, 0, 0, 1.0)
    add_inline(p, "PROBONO · SUMO school-street scenario guidebook", size=9, color=MUTED)
    p_pr = p._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), BORDER)
    border.append(bottom)
    p_pr.append(border)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph_spacing(fp, 0, 0, 1.0)
    page_field(fp)


def cover(doc: Document) -> None:
    for _ in range(5):
        doc.add_paragraph()
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(kicker, 0, 16, 1.0)
    run = kicker.add_run("TECHNICAL GUIDEBOOK")
    set_run(run, size=11, bold=True, color="7A5A00")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(title, 0, 8, 1.0)
    run = title.add_run("SUMO School-Street Scenarios")
    set_run(run, size=30, bold=True, color=DARK_BLUE)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(subtitle, 0, 26, 1.15)
    add_inline(subtitle, "Evidence-aligned baseline, intervention models, sensitivities and publication boundary", size=15, color=BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, 0, 4, 1.0)
    add_inline(p, "PROBONO H2020 · Brussels Living Lab", size=11, color=MUTED)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, 0, 28, 1.0)
    add_inline(p, "De l’Autre Côté de l’École · Auderghem, Brussels", size=11, color=MUTED)

    meta = doc.add_table(rows=5, cols=2)
    fixed_table_geometry(meta, [2700, 6660])
    values = [
        ("Document control", "Value"),
        ("Repository release", "1.1.0"),
        ("Validated model", "v21.2 / SUMO 1.27.1"),
        ("Date", "14 August 2026"),
        ("Purpose", "Technical handover and publication-support record"),
    ]
    for row_index, (row, (label, value)) in enumerate(zip(meta.rows, values)):
        if row_index == 0:
            tr_pr = row._tr.get_or_add_trPr()
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            tr_pr.append(repeat)
        for cell in row.cells:
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), TABLE_FILL if row_index == 0 else LIGHT_FILL)
            cell._tc.get_or_add_tcPr().append(shade)
        row.cells[0].text = ""
        row.cells[1].text = ""
        p0 = row.cells[0].paragraphs[0]
        p1 = row.cells[1].paragraphs[0]
        set_paragraph_spacing(p0, 0, 1, 1.0)
        set_paragraph_spacing(p1, 0, 1, 1.0)
        r0 = p0.add_run(label)
        set_run(r0, size=9.5, bold=True, color=DARK_BLUE)
        r1 = p1.add_run(value)
        set_run(r1, size=9.5, bold=row_index == 0, color=DARK_BLUE if row_index == 0 else INK)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_toc(doc: Document) -> None:
    heading = doc.add_paragraph(style="Heading 1")
    add_inline(heading, "Guide contents", size=16, color=BLUE)
    p = doc.add_paragraph()
    set_paragraph_spacing(p, 0, 8, 1.1)
    add_inline(p, "Page references are fixed for repository release 1.1.0.", size=10, color=MUTED)

    entries = [
        ("Executive summary", "3"),
        ("1. Study question and model boundary", "3"),
        ("2. Evidence hierarchy and interpretation rules", "4"),
        ("3. Baseline construction", "5"),
        ("4. Scenario taxonomy", "8"),
        ("5. Scenario A — enhanced 30 km/h zone", "9"),
        ("6. Scenario B — 30 km/h zone with chicane", "10"),
        ("7. Scenario C0 — shared-space operational proxy", "11"),
        ("8. Scenario C1 — pedestrian-priority sensitivity", "12"),
        ("9. Scenario D — full one-way option", "13"),
        ("10. Scenario E — reduced one-way option", "14"),
        ("11. Demand-response sensitivities", "15"),
        ("12. Tested but not adopted: bounded diversion", "15"),
        ("13. KPI definitions and comparability", "16"),
        ("14. Validation and known warnings", "17"),
        ("15. Reproduction workflow", "18"),
        ("16. Repository packages", "19"),
        ("17. Manuscript use", "19"),
        ("Appendix A. Key files", "20"),
        ("Appendix B. Terminology", "20"),
    ]
    table = doc.add_table(rows=len(entries) + 1, cols=2)
    fixed_table_geometry(table, [8500, 860])
    table.autofit = False
    header = table.rows[0]
    tr_pr = header._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)
    for cell, label in zip(header.cells, ("Section", "Page")):
        cell.text = ""
        shade = OxmlElement("w:shd")
        shade.set(qn("w:fill"), TABLE_FILL)
        cell._tc.get_or_add_tcPr().append(shade)
        p = cell.paragraphs[0]
        set_paragraph_spacing(p, 0, 0, 1.0)
        if label == "Page":
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = p.add_run(label)
        set_run(run, size=9.3, bold=True, color=DARK_BLUE)
    for row, (label, page) in zip(table.rows[1:], entries):
        left, right = row.cells
        left.text = ""
        right.text = ""
        lp = left.paragraphs[0]
        rp = right.paragraphs[0]
        set_paragraph_spacing(lp, 0, 0, 1.0)
        set_paragraph_spacing(rp, 0, 0, 1.0)
        rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        add_inline(lp, label, size=9.3, color=INK)
        add_inline(rp, page, size=9.3, color=MUTED)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def parse_markdown(doc: Document, source: Path, bullet_id: int, number_id: int) -> None:
    lines = source.read_text(encoding="utf-8").splitlines()
    # Skip the Markdown title and metadata block; the DOCX has a designed cover.
    start = next(i for i, line in enumerate(lines) if line.startswith("## Executive summary"))
    lines = lines[start:]
    index = 0
    figure = 0
    in_code = False
    code_lines: list[str] = []

    while index < len(lines):
        line = lines[index]
        if line.startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.22)
                p.paragraph_format.right_indent = Inches(0.12)
                set_paragraph_spacing(p, 4, 7, 1.0)
                paragraph_shading(p, LIGHT_FILL)
                run = p.add_run("\n".join(code_lines))
                set_run(run, size=8.5, color=DARK_BLUE, font="Consolas")
                in_code = False
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        if not line.strip():
            index += 1
            continue
        if line.startswith("## "):
            p = doc.add_paragraph(style="Heading 1")
            add_inline(p, line[3:], size=16, color=BLUE)
            index += 1
            continue
        if line.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            add_inline(p, line[4:], size=13, color=BLUE)
            index += 1
            continue
        if line.startswith("#### "):
            p = doc.add_paragraph(style="Heading 3")
            add_inline(p, line[5:], size=12, color=DARK_BLUE)
            index += 1
            continue
        image_match = re.match(r"!\[([^]]+)\]\(([^)]+)\)", line)
        if image_match:
            figure += 1
            alt, relative = image_match.groups()
            add_image(doc, (source.parent / relative).resolve(), alt, figure)
            index += 1
            continue
        if line.startswith("|") and index + 1 < len(lines) and re.match(r"^\|?\s*:?-+", lines[index + 1]):
            rows = []
            header = [cell.strip() for cell in line.strip().strip("|").split("|")]
            rows.append(header)
            index += 2
            while index < len(lines) and lines[index].startswith("|"):
                rows.append([cell.strip() for cell in lines[index].strip().strip("|").split("|")])
                index += 1
            add_table(doc, rows)
            continue
        if line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.24)
            p.paragraph_format.right_indent = Inches(0.12)
            set_paragraph_spacing(p, 4, 8, 1.2)
            paragraph_shading(p, CAUTION_FILL, "E2C15C")
            add_inline(p, line[2:], size=10.5, color=INK)
            index += 1
            continue
        if re.match(r"^- ", line):
            p = doc.add_paragraph()
            apply_num(p, bullet_id)
            set_paragraph_spacing(p, 0, 4, 1.25)
            add_inline(p, line[2:])
            index += 1
            continue
        if re.match(r"^\d+\. ", line):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.38)
            p.paragraph_format.first_line_indent = Inches(-0.20)
            set_paragraph_spacing(p, 0, 4, 1.25)
            add_inline(p, line)
            index += 1
            continue

        # Join ordinary Markdown lines into one prose paragraph.
        paragraph_lines = [line.strip()]
        index += 1
        while index < len(lines):
            candidate = lines[index]
            if not candidate.strip():
                break
            if candidate.startswith(("#", "- ", "> ", "```", "![", "|")) or re.match(r"^\d+\. ", candidate):
                break
            paragraph_lines.append(candidate.strip())
            index += 1
        p = doc.add_paragraph()
        set_paragraph_spacing(p, 0, 6, 1.25)
        add_inline(p, " ".join(paragraph_lines))


def main() -> None:
    doc = Document()
    doc.core_properties.title = "SUMO School-Street Scenario Guidebook"
    doc.core_properties.subject = "Evidence-aligned baseline, interventions and publication boundary"
    doc.core_properties.author = "Filippos Lygerakis"
    doc.core_properties.keywords = "SUMO, school street, traffic microsimulation, PROBONO"
    setup_styles(doc)
    setup_page(doc.sections[0])
    bullet_id, number_id = numbering(doc)
    cover(doc)
    add_toc(doc)
    parse_markdown(doc, SOURCE, bullet_id, number_id)

    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
